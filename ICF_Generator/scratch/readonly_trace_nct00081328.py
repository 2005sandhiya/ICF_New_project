import sys
sys.path.insert(0, 'backend')
import io
import docx
from pypdf import PdfReader

from app.services.icf_generation_service import icf_generation_service
from app.services.docx_generation_service import docx_generation_service
from app.services.pdf_generation_service import pdf_generation_service
from app.services.template_service import template_service, inspect_docx, TOKEN
from app.services.template_renderer import build_placeholder_values, render_template, populate_heading, replace_tokens, remove_unsupported_sections_and_renumber, prune_subsections_and_empty_blocks

study_id = "NCT00081328"

print("==================================================")
print(f"STEP 1 — CHECK GENERATED ICFDRAFT ({study_id})")
print("==================================================")
from app.models.icf import ICFGenerationRequest
resp = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_id))
draft = resp.icf_draft

sections_to_check = ["purpose", "why_asked", "procedures", "duration", "interventions", "voluntary", "confidentiality", "questions_contact"]
for sec_id in sections_to_check:
    sec = next((s for s in draft.sections if s.section_id == sec_id), None)
    if sec:
        status_str = "NON-EMPTY" if sec.content and sec.content.strip() else "EMPTY"
        length = len(sec.content or '')
        snippet = sec.content[:60].replace('\n', ' ') if sec.content else ''
        print(f"  {sec_id:<20}: {status_str}, {length} chars | '{snippet}'")
    else:
        print(f"  {sec_id:<20}: MISSING FROM DRAFT SECTIONS")

print("\n==================================================")
print("STEP 2 — CHECK PLACEHOLDER VALUES DICTIONARY")
print("==================================================")
facts = icf_generation_service.get_facts(study_id)
values = build_placeholder_values(draft, study_id, "patient_facing", facts=facts)

for key in ["purpose", "voluntary", "procedures", "duration", "confidentiality", "why_asked", "interventions", "questions_contact", "study_title", "STUDY_TITLE", "STUDY_ID", "INVESTIGATOR", "SPONSOR"]:
    val = values.get(key, "")
    status_str = f"VALUE PRESENT ({len(val)} chars)" if val and val.strip() else "EMPTY"
    snippet = val[:60].replace('\n', ' ') if val else ''
    print(f"  {key:<20} -> {status_str} | '{snippet}'")

print("\n==================================================")
print("STEP 3 — CHECK ACTUAL WHO TEMPLATE TOKENS & HEADINGS")
print("==================================================")
template_bytes = template_service.read(study_id, template_service.selection(study_id))
tpl_doc = docx.Document(io.BytesIO(template_bytes))

discovered_tokens = []
for p in tpl_doc.paragraphs:
    for m in TOKEN.finditer(p.text):
        raw = m.group(1) or m.group(2)
        discovered_tokens.append((raw, p.text.strip()))

for tbl in tpl_doc.tables:
    for r in tbl.rows:
        for cell in r.cells:
            for p in cell.paragraphs:
                for m in TOKEN.finditer(p.text):
                    raw = m.group(1) or m.group(2)
                    discovered_tokens.append((raw, p.text.strip()))

print("Discovered Tokens/Placeholders in DOCX:")
for tok, p_text in discovered_tokens:
    print(f"  Token: '{tok}' in paragraph: '{p_text}'")

tpl_doc_obj, slots = inspect_docx(template_bytes)
print("\nDiscovered Slots by inspect_docx():")
for slot_name, (mode, p) in slots.items():
    print(f"  Slot: '{slot_name:<20}' Mode: '{mode:<8}' Paragraph Text: '{p.text}'")

print("\n==================================================")
print("STEP 4 — CHECK RENDERED DOCX PARAGRAPHS")
print("==================================================")
rendered_docx_bytes = render_template(template_bytes, study_id, draft, "patient_facing", facts=facts)
rendered_doc = docx.Document(io.BytesIO(rendered_docx_bytes))

print("Rendered DOCX Paragraphs (Non-empty):")
for idx, p in enumerate(rendered_doc.paragraphs):
    if p.text.strip():
        style_name = p.style.name if p.style else "NoStyle"
        print(f"  P{idx:02d} [{style_name:<12}]: {p.text[:100]}")

print("\nRendered DOCX Tables:")
for idx, tbl in enumerate(rendered_doc.tables):
    print(f"  Table {idx}:")
    for r_idx, r in enumerate(tbl.rows):
        row_text = " | ".join(c.text.strip().replace('\n', ' ') for c in r.cells)
        print(f"    Row {r_idx}: {row_text}")

print("\n==================================================")
print("STEP 5 — CHECK STEP-BY-STEP TRANSFORMATION / DELETION")
print("==================================================")
# Step by step simulation of render_template
step_doc, step_slots = inspect_docx(template_bytes)
print("Step 5.1: After inspect_docx, paragraph count:", len(step_doc.paragraphs))

# 5.2 populate_heading / replace_tokens
sections_map = {sec.section_id: sec for sec in draft.sections}
for name, (mode, paragraph) in step_slots.items():
    if mode == "heading":
        sec = sections_map.get(name)
        if sec and sec.status != "NOT_FOUND" and not (sec.content or "").startswith("Not provided"):
            text_val = values.get(name) or sec.content
            populate_heading(paragraph, text_val)
        else:
            populate_heading(paragraph, "")

print("\nParagraphs after populate_heading:")
for idx, p in enumerate(step_doc.paragraphs):
    if p.text.strip():
        style_name = p.style.name if p.style else "NoStyle"
        print(f"  P{idx:02d} [{style_name:<12}]: {p.text[:100]}")

# 5.3 replace_tokens
for p in [docx.text.paragraph.Paragraph(element, step_doc) for element in step_doc.element.body.iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")]:
    replace_tokens(p, values)

print("\nParagraphs after replace_tokens:")
for idx, p in enumerate(step_doc.paragraphs):
    if p.text.strip():
        style_name = p.style.name if p.style else "NoStyle"
        print(f"  P{idx:02d} [{style_name:<12}]: {p.text[:100]}")

# 5.4 remove_unsupported_sections_and_renumber
remove_unsupported_sections_and_renumber(step_doc, sections_map, values)
print("\nParagraphs after remove_unsupported_sections_and_renumber:")
for idx, p in enumerate(step_doc.paragraphs):
    if p.text.strip():
        style_name = p.style.name if p.style else "NoStyle"
        print(f"  P{idx:02d} [{style_name:<12}]: {p.text[:100]}")

# 5.5 prune_subsections_and_empty_blocks
prune_subsections_and_empty_blocks(step_doc, values)
print("\nParagraphs after prune_subsections_and_empty_blocks:")
for idx, p in enumerate(step_doc.paragraphs):
    if p.text.strip():
        style_name = p.style.name if p.style else "NoStyle"
        print(f"  P{idx:02d} [{style_name:<12}]: {p.text[:100]}")

print("\n==================================================")
print("STEP 6 — CHECK PDF CONVERSION OUTPUT")
print("==================================================")
pdf_bytes = pdf_generation_service.generate_pdf(rendered_docx_bytes, study_id, draft, "patient_facing")
pdf_reader = PdfReader(io.BytesIO(pdf_bytes))
print(f"PDF Page Count: {len(pdf_reader.pages)}")

for p_idx, page in enumerate(pdf_reader.pages):
    print(f"\n--- PDF PAGE {p_idx + 1} ---")
    page_text = page.extract_text()
    for line in page_text.splitlines():
        if line.strip():
            print(f"  {line.strip()}")
