import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pathlib import Path

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def build_who_docx():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("PATIENT-FACING INFORMED CONSENT FORM  |  [STUDY_ID]")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 100, 100)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("WHO-based Informed Consent Form  •  Participant Plain-Language Document")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 120, 120)

    # Document Styles
    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)

    # Header Title Banner
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("INFORMED CONSENT FORM")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(24, 76, 120)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(16)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("PART I: INFORMATION SHEET & PART II: CERTIFICATE OF CONSENT")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(90, 105, 120)

    # Metadata Box Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Protocol Title:", "[STUDY_TITLE]"),
        ("Study ID / NCT Number:", "[STUDY_ID]"),
        ("Study Phase", "[STUDY_PHASE]"),
        ("Principal Investigator:", "[INVESTIGATOR]"),
        ("Research Site / Organization:", "[SITE_NAME_ADDRESS]"),
        ("Sponsor:", "[SPONSOR]")
    ]

    for row_idx, (label, val_token) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_lbl = row.cells[0]
        cell_val = row.cells[1]
        
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        set_cell_background(cell_lbl, "F0 F4 F8")
        set_cell_background(cell_val, "FA FA FA")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)

        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(30, 60, 90)

        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val_token)
        r1.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # PART I HEADER
    h_part1 = doc.add_paragraph(style="Heading 1")
    h_part1.paragraph_format.space_before = Pt(16)
    h_part1.paragraph_format.space_after = Pt(8)
    r_p1 = h_part1.add_run("PART I: INFORMATION SHEET")
    r_p1.font.name = "Arial"
    r_p1.font.size = Pt(14)
    r_p1.font.bold = True
    r_p1.font.color.rgb = RGBColor(24, 76, 120)

    # Section Headings (Heading mode rendering for automatic section removal/renumbering)
    sections = [
        "1. Introduction",
        "2. Purpose of the Research",
        "3. Type of Research Intervention",
        "4. Participant Selection",
        "5. Voluntary Participation",
        "6. Procedures and Protocol",
        "7. Duration",
        "8. Side Effects and Risks",
        "9. Potential Benefits",
        "10. Confidentiality",
        "11. Alternatives to Participating",
        "12. Who to Contact",
    ]

    for heading_text in sections:
        hp = doc.add_paragraph(style="Heading 1")
        hp.paragraph_format.space_before = Pt(14)
        hp.paragraph_format.space_after = Pt(4)
        hp.paragraph_format.keep_with_next = True
        hrun = hp.add_run(heading_text)
        hrun.font.name = "Arial"
        hrun.font.size = Pt(12)
        hrun.font.bold = True
        hrun.font.color.rgb = RGBColor(24, 76, 120)

        # Body paragraph placeholder for heading mode population
        tp = doc.add_paragraph()
        tp.paragraph_format.space_before = Pt(0)
        tp.paragraph_format.space_after = Pt(8)

    # PART II HEADER
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    h_part2 = doc.add_paragraph(style="Heading 1")
    h_part2.paragraph_format.space_before = Pt(20)
    h_part2.paragraph_format.space_after = Pt(8)
    r_p2 = h_part2.add_run("PART II: CERTIFICATE OF CONSENT")
    r_p2.font.name = "Arial"
    r_p2.font.size = Pt(14)
    r_p2.font.bold = True
    r_p2.font.color.rgb = RGBColor(24, 76, 120)

    # Consent Statement
    p_stmt = doc.add_paragraph()
    p_stmt.paragraph_format.space_before = Pt(4)
    p_stmt.paragraph_format.space_after = Pt(12)
    r_stmt = p_stmt.add_run(
        "I have read the foregoing information, or it has been read to me. "
        "I have had the opportunity to ask questions about it and any questions I have asked "
        "have been answered to my satisfaction. I consent voluntarily to be a participant in this study."
    )
    r_stmt.font.bold = True
    r_stmt.font.size = Pt(10.5)

    # Participant Signature Box
    doc.add_paragraph("Participant Consent Signature:").runs[0].font.bold = True
    tbl_part = doc.add_table(rows=3, cols=2)
    tbl_part.alignment = WD_TABLE_ALIGNMENT.CENTER
    part_fields = [
        ("Participant Name:", "[PARTICIPANT_NAME]"),
        ("Participant Signature:", "[PARTICIPANT_SIGNATURE_OR_THUMB_IMPRESSION]"),
        ("Date (Day / Month / Year):", "[PARTICIPANT_SIGN_DATE]")
    ]
    for idx, (label, token) in enumerate(part_fields):
        row = tbl_part.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.5), Inches(4.0)
        set_cell_background(c0, "F8 F9 FA")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        c0.paragraphs[0].add_run(label).font.bold = True
        c1.paragraphs[0].add_run(token)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Witness Section (If illiterate)
    p_wit_hdr = doc.add_paragraph()
    p_wit_hdr.paragraph_format.space_before = Pt(8)
    p_wit_hdr.paragraph_format.space_after = Pt(4)
    r_wit_hdr = p_wit_hdr.add_run("If illiterate (Independent Witness Section):")
    r_wit_hdr.font.bold = True

    p_wit_stmt = doc.add_paragraph()
    p_wit_stmt.paragraph_format.space_after = Pt(8)
    p_wit_stmt.add_run(
        "I have witnessed the accurate reading of the consent form to the potential participant, "
        "and the individual has had the opportunity to ask questions. I confirm that the individual "
        "has given consent freely."
    )

    tbl_wit = doc.add_table(rows=3, cols=2)
    tbl_wit.alignment = WD_TABLE_ALIGNMENT.CENTER
    wit_fields = [
        ("Print Name of Witness:", "[WITNESS_NAME]"),
        ("Signature of Witness:", "[WITNESS_SIGNATURE]"),
        ("Date (Day / Month / Year):", "[WITNESS_SIGN_DATE]")
    ]
    for idx, (label, token) in enumerate(wit_fields):
        row = tbl_wit.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.5), Inches(4.0)
        set_cell_background(c0, "F8 F9 FA")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        c0.paragraphs[0].add_run(label).font.bold = True
        c1.paragraphs[0].add_run(token)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Researcher Statement
    p_res_hdr = doc.add_paragraph()
    p_res_hdr.paragraph_format.space_before = Pt(8)
    p_res_hdr.paragraph_format.space_after = Pt(4)
    p_res_hdr.add_run("Statement by the Researcher / Person Taking Consent:").font.bold = True

    p_res_stmt = doc.add_paragraph()
    p_res_stmt.paragraph_format.space_after = Pt(8)
    p_res_stmt.add_run(
        "I have accurately read out the information sheet to the potential participant, and to the best of my "
        "ability made sure that the participant understands the study procedures, risks, and benefits. I confirm "
        "that the participant was given an opportunity to ask questions about the study, and all questions asked "
        "have been answered correctly and to the best of my ability. A copy of this ICF has been provided to the participant."
    )

    tbl_res = doc.add_table(rows=3, cols=2)
    tbl_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    res_fields = [
        ("Print Name of Researcher:", "[PERSON_OBTAINING_CONSENT_NAME]"),
        ("Signature of Researcher:", "[PERSON_OBTAINING_CONSENT_SIGNATURE]"),
        ("Date (Day / Month / Year):", "[CONSENT_DATE]")
    ]
    for idx, (label, token) in enumerate(res_fields):
        row = tbl_res.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width, c1.width = Inches(2.5), Inches(4.0)
        set_cell_background(c0, "F8 F9 FA")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        c0.paragraphs[0].add_run(label).font.bold = True
        c1.paragraphs[0].add_run(token)

    output_path = Path("data/templates/ICF_Clinical_Study_Template_FINAL.docx")
    doc.save(output_path)
    print(f"Successfully created WHO-based template at {output_path}")

if __name__ == "__main__":
    build_who_docx()
