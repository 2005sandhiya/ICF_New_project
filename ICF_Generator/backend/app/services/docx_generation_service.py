import io
import re
from datetime import datetime, timezone
import logging
from typing import Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from app.models.icf import ICFDraft, FactStatusEnum

logger = logging.getLogger("icf_generator.docx_generation_service")


class DOCXGenerationService:
    """
    Phase 6 Part 5: Service for generating formatted Word (.docx) documents
    from structured ICF drafts and patient-friendly transformations.
    """

    def generate_docx(
        self,
        study_id: str,
        icf_draft: ICFDraft,
        export_type: str = "standard",
        study_title: Optional[str] = None,
        review_status: Optional[str] = "APPROVED",
        template_content: Optional[bytes] = None,
    ) -> bytes:
        """
        Generates a Word (.docx) binary for the provided ICF draft.

        Args:
            study_id: Target study identifier.
            icf_draft: Structured ICFDraft to export.
            export_type: "standard" or "patient_friendly".
            study_title: Optional study title text.
            review_status: Current human review status tag.
        """
        is_custom = False
        if template_content is None:
            from app.services.template_service import template_service, DEFAULT_TEMPLATE_PATH, TemplateError
            try:
                selected = template_service.selection(study_id)
                template_content = template_service.read(study_id, selected)
                is_custom = (selected.kind == "custom")
            except Exception:
                template_content = None
            if template_content is None and DEFAULT_TEMPLATE_PATH.is_file():
                template_content = DEFAULT_TEMPLATE_PATH.read_bytes()
        else:
            from app.services.template_service import template_service
            try:
                selected = template_service.selection(study_id)
                is_custom = (selected.kind == "custom")
            except Exception:
                is_custom = False

        if template_content is not None:
            from app.services.template_renderer import render_template
            return render_template(template_content, study_id, icf_draft, export_type, is_custom=is_custom)
        
        from app.services.template_service import TemplateError, DEFAULT_TEMPLATE_PATH
        raise TemplateError(f"Production ICF template is missing or invalid. Ensure {DEFAULT_TEMPLATE_PATH} exists.")
        doc = docx.Document()
        doc.core_properties.title = f"{study_id} {export_type} Revision {icf_draft.revision}"
        doc.core_properties.subject = f"Study {study_id}; Revision {icf_draft.revision}; {export_type}"

        # Set 1-inch margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        doc_type_label = "PATIENT-FACING INFORMED CONSENT FORM (ICF)"

        # Main Title Header
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_title.paragraph_format.space_before = Pt(0)
        p_title.paragraph_format.space_after = Pt(2)
        run_title = p_title.add_run(doc_type_label)
        run_title.bold = True
        run_title.font.size = Pt(18)
        run_title.font.color.rgb = RGBColor(15, 23, 42)

        # Subtitle
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(14)
        run_sub = p_sub.add_run("Participant Plain-Language Consent Document")
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
        run_sub.font.color.rgb = RGBColor(100, 116, 139)

        # Metadata Table
        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Extract study title from title section if not provided explicitly
        if not study_title:
            title_sec = next(
                (
                    s
                    for s in icf_draft.sections
                    if s.section_id in ["study_title", "title"]
                ),
                None,
            )
            if (
                title_sec
                and title_sec.content
                and not title_sec.content.startswith("Not provided in the study protocol.")
            ):
                study_title = title_sec.content
            else:
                study_title = f"Clinical Research Protocol {study_id}"

        now_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        status_display = f"{review_status or 'APPROVED'} (Internal Review)"

        metadata = [
            ("Target Study ID:", study_id),
            ("Study Title:", study_title),
            ("Document Revision:", f"Revision {icf_draft.revision}"),
            ("Human Review Status:", status_display),
            ("Export Date:", now_str),
        ]

        for i, (label, val) in enumerate(metadata):
            row = table.rows[i]
            cell_lbl = row.cells[0]
            cell_val = row.cells[1]

            cell_lbl.paragraphs[0].paragraph_format.space_after = Pt(2)
            run_l = cell_lbl.paragraphs[0].add_run(label)
            run_l.bold = True
            run_l.font.size = Pt(9.5)
            run_l.font.color.rgb = RGBColor(51, 65, 85)

            cell_val.paragraphs[0].paragraph_format.space_after = Pt(2)
            run_v = cell_val.paragraphs[0].add_run(val)
            run_v.font.size = Pt(9.5)
            run_v.font.color.rgb = RGBColor(15, 23, 42)
            if label == "Human Review Status:":
                run_v.bold = True
                run_v.font.color.rgb = RGBColor(16, 185, 129)

        # Disclaimer Box
        p_disc = doc.add_paragraph()
        p_disc.paragraph_format.space_before = Pt(14)
        p_disc.paragraph_format.space_after = Pt(16)
        run_disc_head = p_disc.add_run("INTERNAL WORKFLOW DISCLAIMER:\n")
        run_disc_head.bold = True
        run_disc_head.font.size = Pt(8.5)
        run_disc_head.font.color.rgb = RGBColor(180, 83, 9)

        run_disc = p_disc.add_run(
            "This document is an AI-generated Informed Consent Form (ICF) draft based on protocol evidence facts. "
            f"Human Review status for this revision is {review_status} for internal clinical preparation. "
            "This document is intended for clinical review and protocol development and does NOT constitute regulatory or IRB approval."
        )
        run_disc.font.size = Pt(8.5)
        run_disc.font.italic = True
        run_disc.font.color.rgb = RGBColor(100, 116, 139)

        # Render ICF Sections
        sec_num = 1
        for sec in icf_draft.sections:
            if (
                sec.status == FactStatusEnum.NOT_FOUND
                or not sec.content
                or sec.content.strip().startswith("Not provided in the study protocol.")
            ):
                continue

            p_sec_head = doc.add_heading(level=2)
            p_sec_head.paragraph_format.space_before = Pt(12)
            p_sec_head.paragraph_format.space_after = Pt(4)
            # Strip existing leading numbers if present before adding sequential sec_num
            clean_title = re.sub(r"^\s*\d+\.\s*", "", sec.title)
            run_head = p_sec_head.add_run(f"{sec_num}. {clean_title}")
            sec_num += 1
            run_head.font.size = Pt(13)
            run_head.font.bold = True
            run_head.font.color.rgb = RGBColor(30, 41, 59)

            # Section Content
            p_content = doc.add_paragraph()
            p_content.paragraph_format.space_after = Pt(6)
            p_content.paragraph_format.line_spacing = 1.15

            run_c = p_content.add_run(sec.content)
            run_c.font.size = Pt(10.5)
            run_c.font.color.rgb = RGBColor(30, 41, 59)

            # Source evidence line if available
            if sec.source_evidence_ids and len(sec.source_evidence_ids) > 0:
                p_ev = doc.add_paragraph()
                p_ev.paragraph_format.space_after = Pt(8)
                run_ev = p_ev.add_run(
                    f"Supporting Evidence Chunks: {', '.join(sec.source_evidence_ids)}"
                )
                run_ev.font.size = Pt(8.5)
                run_ev.font.italic = True
                run_ev.font.color.rgb = RGBColor(148, 163, 184)

        doc.add_page_break()

        # Signature Block Header
        p_sig_head = doc.add_heading(level=2)
        p_sig_head.paragraph_format.space_before = Pt(12)
        p_sig_head.paragraph_format.space_after = Pt(6)
        run_sig_head = p_sig_head.add_run(
            "PARTICIPANT CONSENT ACKNOWLEDGMENT & SIGNATURES"
        )
        run_sig_head.font.size = Pt(13)
        run_sig_head.font.bold = True
        run_sig_head.font.color.rgb = RGBColor(30, 41, 59)

        p_sig_intro = doc.add_paragraph()
        p_sig_intro.paragraph_format.space_after = Pt(14)
        run_sig_intro = p_sig_intro.add_run(
            "By signing below, I acknowledge that I have read and understood the information provided in this consent document. "
            "I have had the opportunity to ask questions and all of my questions have been answered to my satisfaction. "
            "I voluntarily agree to participate in this study."
        )
        run_sig_intro.font.size = Pt(10)
        run_sig_intro.font.color.rgb = RGBColor(51, 65, 85)

        # Lines for signatures
        p_lines1 = doc.add_paragraph()
        p_lines1.paragraph_format.space_before = Pt(10)
        p_lines1.paragraph_format.space_after = Pt(2)
        p_lines1.add_run(
            "___________________________________________              ______________________"
        )

        p_lbl1 = doc.add_paragraph()
        p_lbl1.paragraph_format.space_after = Pt(14)
        run_l1 = p_lbl1.add_run(
            "Participant Name (Printed)                                                    Date"
        )
        run_l1.font.size = Pt(9)
        run_l1.font.color.rgb = RGBColor(100, 116, 139)

        p_lines2 = doc.add_paragraph()
        p_lines2.paragraph_format.space_after = Pt(2)
        p_lines2.add_run(
            "___________________________________________              ______________________"
        )

        p_lbl2 = doc.add_paragraph()
        p_lbl2.paragraph_format.space_after = Pt(20)
        run_l2 = p_lbl2.add_run(
            "Participant Signature                                                         Date"
        )
        run_l2.font.size = Pt(9)
        run_l2.font.color.rgb = RGBColor(100, 116, 139)

        p_lines3 = doc.add_paragraph()
        p_lines3.paragraph_format.space_after = Pt(2)
        p_lines3.add_run(
            "___________________________________________              ______________________"
        )

        p_lbl3 = doc.add_paragraph()
        p_lbl3.paragraph_format.space_after = Pt(14)
        run_l3 = p_lbl3.add_run(
            "Person Obtaining Consent Name (Printed)                                        Date"
        )
        run_l3.font.size = Pt(9)
        run_l3.font.color.rgb = RGBColor(100, 116, 139)

        p_lines4 = doc.add_paragraph()
        p_lines4.paragraph_format.space_after = Pt(2)
        p_lines4.add_run(
            "___________________________________________              ______________________"
        )

        p_lbl4 = doc.add_paragraph()
        run_l4 = p_lbl4.add_run(
            "Person Obtaining Consent Signature                                             Date"
        )
        run_l4.font.size = Pt(9)
        run_l4.font.color.rgb = RGBColor(100, 116, 139)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


docx_generation_service = DOCXGenerationService()
