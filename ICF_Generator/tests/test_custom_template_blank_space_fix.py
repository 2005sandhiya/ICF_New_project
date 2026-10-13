import io
import os
import sys
import pytest
from docx import Document

sys.path.insert(0, os.path.abspath("backend"))

from app.models.icf import ICFDraft, ICFSection
from app.models.protocol_facts import ProtocolFacts, TraceableFact, FactStatusEnum
from app.services.template_renderer import render_template, document_text
from app.services.template_service import DEFAULT_TEMPLATE_PATH

def create_template_with_empty_spacers():
    doc = Document()
    doc.add_heading("INFORMED CONSENT FORM", level=1)
    doc.add_paragraph("[STUDY_TITLE]")
    
    # Add empty spacer paragraphs that used to cause large blank gaps
    doc.add_paragraph("")
    doc.add_paragraph("")
    
    doc.add_heading("1. Purpose of the Study", level=2)
    doc.add_paragraph("[PURPOSE_OF_STUDY]")
    
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_no_empty_paragraphs_between_header_and_first_section():
    template_bytes = create_template_with_empty_spacers()
    draft = ICFDraft(sections=[
        ICFSection(section_id="study_title", title="Study Title", content="BioX Study 2026", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Evaluate safety and efficacy.", status=FactStatusEnum.SUPPORTED),
    ])
    facts = ProtocolFacts(study_id="BLANK-GAP-001")
    rendered = render_template(template_bytes, "BLANK-GAP-001", draft, "patient_facing", facts=facts)
    rendered_doc = Document(io.BytesIO(rendered))
    
    # Extract body paragraphs
    non_empty = [p.text.strip() for p in rendered_doc.paragraphs if p.text.strip()]
    empty_paras = [p for p in rendered_doc.paragraphs if not p.text.strip()]
    
    # Verify empty spacer paragraphs were pruned
    assert len(empty_paras) == 0
    assert any("INFORMED CONSENT FORM" in t for t in non_empty)
    assert any("BioX Study 2026" in t for t in non_empty)
    assert any("Purpose of the Study" in t for t in non_empty)
    assert any("Evaluate safety and efficacy." in t for t in non_empty)

def test_meaningful_formatting_and_headings_preserved():
    path = os.path.join("data/templates/custom_demo", "ICF_Custom_Template_A_Simple.docx")
    assert os.path.exists(path)
    with open(path, "rb") as f:
        template_bytes = f.read()
    
    draft = ICFDraft(sections=[
        ICFSection(section_id="study_title", title="Study Title", content="BioX Simple Study", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Simple purpose text.", status=FactStatusEnum.SUPPORTED),
    ])
    facts = ProtocolFacts(study_id="FORMAT-TEST")
    rendered = render_template(template_bytes, "FORMAT-TEST", draft, "patient_facing", facts=facts)
    rendered_doc = Document(io.BytesIO(rendered))
    
    # Verify headings and section structure remain intact
    headings = [p.text for p in rendered_doc.paragraphs if p.text.startswith("1. Invitation") or p.text.startswith("2. Purpose")]
    assert len(headings) >= 2
    assert "BioX Simple Study" in document_text(rendered)

def test_default_production_template_preserves_content():
    assert DEFAULT_TEMPLATE_PATH.is_file()
    template_bytes = DEFAULT_TEMPLATE_PATH.read_bytes()
    draft = ICFDraft(sections=[
        ICFSection(section_id="study_title", title="Study Title", content="Default Production Study", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Production purpose text.", status=FactStatusEnum.SUPPORTED),
    ])
    facts = ProtocolFacts(
        study_id="PROD-TEST-2",
        study_title=TraceableFact(value="Default Production Study", status=FactStatusEnum.SUPPORTED)
    )
    rendered = render_template(template_bytes, "PROD-TEST-2", draft, "patient_facing", facts=facts)
    assert len(rendered) > 0
    full_text = document_text(rendered)
    assert "Default Production Study" in full_text
