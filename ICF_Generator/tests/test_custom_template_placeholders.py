import io
import docx
import pytest
from app.services.template_service import template_service, inspect_docx, TemplateError
from app.services.template_renderer import render_template
from app.models.icf import ICFDraft, ICFSection, FactStatusEnum


def create_docx_with_text(paragraphs_list):
    doc = docx.Document()
    for p_text in paragraphs_list:
        doc.add_paragraph(p_text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_1_ordinary_form_fields_accepted():
    # [Name], [Date], [Signature]
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent for a Clinical Study",
        "1. Purpose of the Study",
        "This study evaluates safety.",
        "Participant Name: [Name]",
        "Date of Consent: [Date]",
        "Signature: [Signature]"
    ])
    doc, slots = inspect_docx(docx_bytes)
    assert doc is not None
    assert "purpose" in slots


def test_2_descriptive_form_fields_accepted():
    # [Investigator Name], [Participant Name], [Study Date]
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent for a Clinical Study",
        "Purpose of the Study",
        "Study description text.",
        "Investigator: [Investigator Name]",
        "Participant: [Participant Name]",
        "Date: [Study Date]"
    ])
    doc, slots = inspect_docx(docx_bytes)
    assert doc is not None
    assert "purpose" in slots


def test_3_valid_application_placeholders_recognized():
    # [STUDY_TITLE], [PURPOSE_OF_STUDY]
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent for a Clinical Study",
        "Title: [STUDY_TITLE]",
        "Purpose: [PURPOSE_OF_STUDY]"
    ])
    doc, slots = inspect_docx(docx_bytes)
    assert doc is not None
    assert "purpose" in slots


def test_4_mixed_application_and_ordinary_form_fields():
    # Both valid application placeholders and ordinary form fields
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent for a Clinical Study",
        "Study Title: [STUDY_TITLE]",
        "Purpose of the Study: [PURPOSE_OF_STUDY]",
        "Participant Name: [Name]",
        "Signature: [Signature]"
    ])
    doc, slots = inspect_docx(docx_bytes)
    assert doc is not None
    assert "purpose" in slots

    draft = ICFDraft(sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Grounded purpose content", status=FactStatusEnum.SUPPORTED)
    ])
    rendered = render_template(docx_bytes, "TEST_STUDY", draft, "standard")
    doc_rendered = docx.Document(io.BytesIO(rendered))
    full_text = "\n".join(p.text for p in doc_rendered.paragraphs)
    
    assert "[Name]" in full_text
    assert "[Signature]" in full_text
    assert "Grounded purpose content" in full_text


def test_5_arbitrary_custom_section_titles_accepted():
    # [STUDY_OVERVIEW], {{ STUDY_OVERVIEW }}, [Study Design], [Eligibility], [Costs and Payments]
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent for a Clinical Study",
        "Overview: [STUDY_OVERVIEW]",
        "Summary: {{ STUDY_OVERVIEW }}",
        "1. Purpose of the Study",
        "[PURPOSE_OF_STUDY]",
        "Study Design: [Study Design]",
        "Eligibility: [Eligibility]",
        "Costs: [Costs and Payments]"
    ])
    doc, slots = inspect_docx(docx_bytes)
    assert doc is not None
    assert "purpose" in slots


def test_6_custom_test_templates():
    # Build ICF_Test_Template_A_Concise.docx
    doc_a = docx.Document()
    doc_a.add_heading("Participant Information and Consent Form", level=1)
    doc_a.add_paragraph("Study Title: [STUDY_TITLE]")
    doc_a.add_heading("1. Purpose of the Study", level=1)
    doc_a.add_paragraph("[PURPOSE_OF_STUDY]")
    doc_a.add_paragraph("Participant Name: [Name]")
    doc_a.add_paragraph("Date: [Date]")
    doc_a.add_paragraph("Signature: [Signature]")
    buf_a = io.BytesIO()
    doc_a.save(buf_a)
    bytes_a = buf_a.getvalue()

    sel_a = template_service.upload("TEST_STUDY_A", "ICF_Test_Template_A_Concise.docx", bytes_a)
    assert sel_a.kind == "custom"
    assert not any("Unsupported template placeholder: Name" in w for w in sel_a.warnings)

    # Build ICF_Test_Template_B_Detailed.docx
    doc_b = docx.Document()
    doc_b.add_heading("Informed Consent Document", level=1)
    doc_b.add_paragraph("Protocol ID: [STUDY_ID]")
    doc_b.add_heading("1. Purpose of the Research", level=1)
    doc_b.add_paragraph("[PURPOSE_OF_STUDY]")
    doc_b.add_paragraph("Investigator: [Investigator Name]")
    doc_b.add_paragraph("Participant Signature: [SIGNATURE_AND_DATE]")
    buf_b = io.BytesIO()
    doc_b.save(buf_b)
    bytes_b = buf_b.getvalue()

    sel_b = template_service.upload("TEST_STUDY_B", "ICF_Test_Template_B_Detailed.docx", bytes_b)
    assert sel_b.kind == "custom"
    assert not any("Unsupported template placeholder: Investigator Name" in w for w in sel_b.warnings)


def test_7_custom_template_content_injection():
    # Verify custom template content injection and form field preservation
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent Form",
        "Title: [STUDY_TITLE]",
        "Overview: [STUDY_OVERVIEW]",
        "Design: [STUDY_DESIGN]",
        "Eligibility Rules: [ELIGIBILITY]",
        "Costs and Payment Info: [COSTS_AND_PAYMENT]",
        "New Info: [NEW_INFORMATION]",
        "Post Study Care: [POST_STUDY_CARE]",
        "Participant Name: [Name]",
        "Date of Consent: [Date]",
        "Signature: [Signature]"
    ])

    draft = ICFDraft(sections=[
        ICFSection(section_id="study_title", title="Study Title", content="Grounded Study Alpha Title", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Grounded purpose overview text", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="interventions", title="Study Interventions", content="Grounded treatment arm design details", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="why_asked", title="Why You Are Being Asked", content="Grounded eligibility inclusion criteria details", status=FactStatusEnum.SUPPORTED),
    ])

    rendered = render_template(docx_bytes, "NCT_TEST_ALPHA", draft, "patient_facing")
    doc_rendered = docx.Document(io.BytesIO(rendered))
    full_text = "\n".join(p.text for p in doc_rendered.paragraphs)

    # 1. [STUDY_OVERVIEW] replaced
    assert "[STUDY_OVERVIEW]" not in full_text
    assert "Grounded Study Alpha Title" in full_text
    # 2. [STUDY_DESIGN] replaced
    assert "[STUDY_DESIGN]" not in full_text
    assert "Grounded treatment arm design details" in full_text
    # 3. [ELIGIBILITY] replaced
    assert "[ELIGIBILITY]" not in full_text
    assert "Grounded eligibility inclusion criteria details" in full_text
    # 4. [COSTS_AND_PAYMENT] replaced with grounded fallback
    assert "[COSTS_AND_PAYMENT]" not in full_text
    assert "payment or compensation" in full_text or "no cost" in full_text
    # 5. [NEW_INFORMATION] replaced
    assert "[NEW_INFORMATION]" not in full_text
    assert "new information" in full_text
    # 6. [POST_STUDY_CARE] replaced
    assert "[POST_STUDY_CARE]" not in full_text
    assert "completion of the study" in full_text
    # 7, 8, 9. Ordinary form fields remain intact
    assert "[Name]" in full_text
    assert "[Date]" in full_text
    assert "[Signature]" in full_text
    # 10. Canonical placeholders work
    assert "Grounded Study Alpha Title" in full_text


def test_8_post_study_care_grounding():
    # 1. Test grounded post-trial care evidence present
    docx_bytes = create_docx_with_text([
        "Participant Consent Document",
        "1. Purpose of the Study",
        "[PURPOSE_OF_STUDY]",
        "Post Study Care Information: [POST_STUDY_CARE]"
    ])
    draft_with_care = ICFDraft(sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Sample purpose text", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="post_trial_care", title="Post-Trial Access and Care", content="Participants will receive 12 months of extended safety monitoring care at no cost.", status=FactStatusEnum.SUPPORTED)
    ])
    rendered = render_template(docx_bytes, "NCT_CARE_TEST", draft_with_care, "patient_facing")
    doc_rendered = docx.Document(io.BytesIO(rendered))
    full_text = "\n".join(p.text for p in doc_rendered.paragraphs)

    assert "[POST_STUDY_CARE]" not in full_text
    assert "Participants will receive 12 months of extended safety monitoring care at no cost." in full_text

    # 2. Test absent post-trial care evidence falls back to safe standard disclosure
    draft_no_care = ICFDraft(sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Sample purpose text", status=FactStatusEnum.SUPPORTED)
    ])
    rendered_fallback = render_template(docx_bytes, "NCT_CARE_TEST_2", draft_no_care, "patient_facing")
    doc_fallback = docx.Document(io.BytesIO(rendered_fallback))
    fallback_text = "\n".join(p.text for p in doc_fallback.paragraphs)

    assert "[POST_STUDY_CARE]" not in fallback_text
    assert "Care after completion of the study as described in the protocol." in fallback_text


def test_9_new_custom_placeholder_aliases():
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent Document",
        "1. Purpose of the Study",
        "[PURPOSE_OF_STUDY]",
        "Why Asked: [WHY_ASKED]",
        "Eligibility: [ELIGIBILITY_SUMMARY]",
        "Interventions: [INTERVENTIONS]",
        "Procedures: [PROCEDURES]",
        "Alternatives: [ALTERNATIVES]",
        "Confidentiality: [CONFIDENTIALITY]",
        "Voluntary: [VOLUNTARY]",
        "Registration: [STUDY_REGISTRATION_AND_RESULTS]",
        "Contact: [CONTACT_INFORMATION]",
        "Injury Care: [COMPENSATION_FOR_STUDY_RELATED_INJURY]",
        "Visit 1 Notes: [VISIT_1_NOTES]",
        "Visit 2 Notes: [VISIT_2_NOTES]",
        "Follow Up Notes: [FOLLOW_UP_NOTES]",
        "Participant Name: [NAME]",
        "Date of Consent: [DATE]",
        "Signature: [SIGNATURE]"
    ])

    draft = ICFDraft(sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Grounded purpose content", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="why_asked", title="Why You Are Being Asked", content="Grounded eligibility criteria details", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="interventions", title="Study Interventions", content="Grounded treatment arm details", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="procedures", title="Study Procedures", content="Grounded visit schedule and testing details", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="alternatives", title="Alternatives to Participation", content="Grounded standard care alternatives", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="confidentiality", title="Confidentiality", content="Grounded privacy protection details", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="voluntary", title="Voluntary Participation", content="Grounded voluntary withdrawal details", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="questions_contact", title="Questions / Contact Information", content="Contact Dr. Investigator at 555-0199", status=FactStatusEnum.SUPPORTED),
    ])

    rendered = render_template(docx_bytes, "NCT_ALIAS_TEST", draft, "patient_facing")
    doc_rendered = docx.Document(io.BytesIO(rendered))
    full_text = "\n".join(p.text for p in doc_rendered.paragraphs)

    # 1. Alias replacements
    assert "[WHY_ASKED]" not in full_text
    assert "Grounded eligibility criteria details" in full_text

    assert "[ELIGIBILITY_SUMMARY]" not in full_text

    assert "[INTERVENTIONS]" not in full_text
    assert "Grounded treatment arm details" in full_text

    assert "[PROCEDURES]" not in full_text
    assert "Grounded visit schedule and testing details" in full_text

    assert "[ALTERNATIVES]" not in full_text
    assert "Grounded standard care alternatives" in full_text

    assert "[CONFIDENTIALITY]" not in full_text
    assert "Grounded privacy protection details" in full_text

    assert "[VOLUNTARY]" not in full_text
    assert "Grounded voluntary withdrawal details" in full_text

    assert "[STUDY_REGISTRATION_AND_RESULTS]" not in full_text
    assert "study results sharing" in full_text or "registered on a public trial registry" in full_text

    assert "[CONTACT_INFORMATION]" not in full_text
    assert "Contact Dr. Investigator at 555-0199" in full_text

    assert "[COMPENSATION_FOR_STUDY_RELATED_INJURY]" not in full_text
    assert "Medical treatment will be provided for study-related injuries" in full_text

    # 2. Form fields preserved
    assert "[NAME]" in full_text
    assert "[DATE]" in full_text
    assert "[SIGNATURE]" in full_text

    # 3. Unsupported visit-note placeholders remain unmapped/uninvented
    assert "[VISIT_1_NOTES]" in full_text
    assert "[VISIT_2_NOTES]" in full_text
    assert "[FOLLOW_UP_NOTES]" in full_text


def test_10_visit_schedule_activity_aliases():
    docx_bytes = create_docx_with_text([
        "Participant Information and Consent Document",
        "1. Purpose of the Study",
        "[PURPOSE_OF_STUDY]",
        "Visit 1 Activities: [VISIT_1_ACTIVITIES]",
        "Visit 2 Activities: [VISIT_2_ACTIVITIES]",
        "Follow Up Timing: [FOLLOW_UP_TIMING]",
        "Follow Up Activities: [FOLLOW_UP_ACTIVITIES]",
        "Visit 1 Notes: [VISIT_1_NOTES]",
        "Visit 2 Notes: [VISIT_2_NOTES]",
        "Follow Up Notes: [FOLLOW_UP_NOTES]",
        "Participant Name: [NAME]",
        "Date of Consent: [DATE]",
        "Signature: [SIGNATURE]"
    ])

    draft = ICFDraft(sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Grounded purpose content", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="procedures", title="Study Procedures", content="Grounded procedures and visit schedule content", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="duration", title="Study Duration", content="12 weeks of active treatment", status=FactStatusEnum.SUPPORTED),
    ])

    rendered = render_template(docx_bytes, "NCT_VISIT_TEST", draft, "patient_facing")
    doc_rendered = docx.Document(io.BytesIO(rendered))
    full_text = "\n".join(p.text for p in doc_rendered.paragraphs)

    # A. [VISIT_1_ACTIVITIES] receives grounded content
    assert "[VISIT_1_ACTIVITIES]" not in full_text
    assert "Informed consent, eligibility check" in full_text or "Grounded procedures" in full_text

    # B. [VISIT_2_ACTIVITIES] receives grounded content
    assert "[VISIT_2_ACTIVITIES]" not in full_text
    assert "Study treatment administration" in full_text or "Grounded procedures" in full_text

    # C. [FOLLOW_UP_TIMING] receives grounded duration/timing content
    assert "[FOLLOW_UP_TIMING]" not in full_text
    assert "12 weeks of active treatment" in full_text or "Follow-up period" in full_text

    # D. [FOLLOW_UP_ACTIVITIES] receives grounded follow-up procedure content
    assert "[FOLLOW_UP_ACTIVITIES]" not in full_text
    assert "End-of-study assessments" in full_text or "Grounded procedures" in full_text

    # E. [VISIT_1_NOTES] remains unchanged
    assert "[VISIT_1_NOTES]" in full_text

    # F. [VISIT_2_NOTES] remains unchanged
    assert "[VISIT_2_NOTES]" in full_text

    # G. [FOLLOW_UP_NOTES] remains unchanged
    assert "[FOLLOW_UP_NOTES]" in full_text

    # H. [NAME], [DATE], [SIGNATURE] remain unchanged
    assert "[NAME]" in full_text
    assert "[DATE]" in full_text
    assert "[SIGNATURE]" in full_text


def test_11_risk_aliases_and_fallback_behavior():
    # 1. Test grounded risks render correctly and aliases [RISKS], [RISKS_AND_SIDE_EFFECTS], [SIDE_EFFECTS] map to grounded content
    docx_bytes_grounded = create_docx_with_text([
        "Participant Information and Consent Document",
        "1. Purpose of the Study",
        "[PURPOSE_OF_STUDY]",
        "Risks: [RISKS]",
        "Side Effects: [RISKS_AND_SIDE_EFFECTS]",
        "Other Side Effects: [SIDE_EFFECTS]",
        "Procedure Risks: [PROCEDURE_RISKS]",
        "Common Risks: [IP_RISKS_COMMON]",
        "Participant Name: [Name]",
        "Date: [Date]",
        "Signature: [Signature]"
    ])

    draft_grounded = ICFDraft(sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Grounded study purpose text", status=FactStatusEnum.SUPPORTED),
        ICFSection(section_id="risks", title="Risks and Discomforts", content="Mild headaches, transient fatigue, and localized skin redness.", status=FactStatusEnum.SUPPORTED)
    ])

    rendered_grounded = render_template(docx_bytes_grounded, "TEST_RISK_GROUNDED", draft_grounded, "patient_facing")
    doc_g = docx.Document(io.BytesIO(rendered_grounded))
    full_text_g = "\n".join(p.text for p in doc_g.paragraphs)

    assert "[RISKS]" not in full_text_g
    assert "[RISKS_AND_SIDE_EFFECTS]" not in full_text_g
    assert "[SIDE_EFFECTS]" not in full_text_g
    assert "Mild headaches, transient fatigue, and localized skin redness." in full_text_g
    assert "[Name]" in full_text_g
    assert "[Date]" in full_text_g
    assert "[Signature]" in full_text_g

    # 2. Test absent grounded risks produce exact evidence-absence statement and do not leave Section 9 blank
    docx_bytes_absent = create_docx_with_text([
        "Participant Information and Consent Document",
        "1. Purpose of the Study",
        "[PURPOSE_OF_STUDY]",
        "9. Risks and Side Effects",
        "[RISKS]",
        "Procedure Risks: [PROCEDURE_RISKS]",
        "Common Risks: [IP_RISKS_COMMON]",
        "Unknown Risks: [UNKNOWN_RISKS_STATEMENT]",
        "Participant Name: [Name]",
        "Date: [Date]",
        "Signature: [Signature]"
    ])

    draft_absent = ICFDraft(sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content="Grounded study purpose text", status=FactStatusEnum.SUPPORTED)
    ])

    rendered_absent = render_template(docx_bytes_absent, "TEST_RISK_ABSENT", draft_absent, "patient_facing")
    doc_a = docx.Document(io.BytesIO(rendered_absent))
    full_text_a = "\n".join(p.text for p in doc_a.paragraphs)

    # A. Section 9 heading is retained when risks are absent
    assert "9. Risks and Side Effects" in full_text_a

    # B. Placeholders are populated with evidence-absence statements
    assert "[RISKS]" not in full_text_a
    assert "not provided in the study protocol" in full_text_a

    # C. Form fields remain untouched
    assert "[Name]" in full_text_a
    assert "[Date]" in full_text_a
    assert "[Signature]" in full_text_a






