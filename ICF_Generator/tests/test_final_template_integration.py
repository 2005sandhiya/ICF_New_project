import io
from pathlib import Path
import docx
import pytest

from app.models.icf import ICFDraft, ICFSection, FactStatusEnum
from app.services.docx_generation_service import docx_generation_service
from app.services.template_renderer import document_text, render_template
from app.services.template_service import DEFAULT_TEMPLATE_PATH, template_service


def create_sample_draft(study_id="NCT00081328", revision=1):
    return ICFDraft(
        study_id=study_id,
        revision=revision,
        sections=[
            ICFSection(section_id="study_title", title="Study Title", content="Treatment Options for Type 2 Diabetes in Adolescents and Youth (TODAY)", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e1"]),
            ICFSection(section_id="invitation", title="Invitation to Participate", content="You are invited to take part in the TODAY clinical trial.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e2"]),
            ICFSection(section_id="purpose", title="Purpose of the Study", content="The purpose of this study is to compare three treatment safety approaches.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e3"]),
            ICFSection(section_id="why_asked", title="Why You Are Being Asked", content="You are invited because you are between 10 and 17 years old with recently diagnosed type 2 diabetes.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e4"]),
            ICFSection(section_id="procedures", title="Study Procedures", content="You will have regular clinic visits every 2 to 3 months and physical exams.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e5"]),
            ICFSection(section_id="duration", title="Study Duration", content="Your participation in the study will last for up to 5 years.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e6"]),
            ICFSection(section_id="interventions", title="Study Interventions / Treatments", content="Participants are assigned to metformin alone, metformin plus rosiglitazone, or metformin plus lifestyle intervention.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e7"]),
            ICFSection(section_id="risks", title="Risks and Discomforts", content="Possible side effects include stomach upset, hypoglycemia, and blood draw discomfort.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e8"]),
            ICFSection(section_id="benefits", title="Potential Benefits", content="You may experience improved blood sugar control.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e9"]),
            ICFSection(section_id="alternatives", title="Alternatives to Participation", content="Not provided in the study protocol.", status=FactStatusEnum.NOT_FOUND),
            ICFSection(section_id="voluntary", title="Voluntary Participation", content="Taking part in this study is completely voluntary.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e10"]),
            ICFSection(section_id="confidentiality", title="Confidentiality", content="Your research records will be kept confidential to the extent permitted by law.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e11"]),
            ICFSection(section_id="questions_contact", title="Questions / Contact Information", content="Contact the Principal Investigator at 555-0199 for study questions.", status=FactStatusEnum.SUPPORTED, source_evidence_ids=["e12"]),
        ]
    )


def test_builtin_template_file_exists():
    assert DEFAULT_TEMPLATE_PATH.is_file()
    selected = template_service.selection("STUDY_TEST")
    content = template_service.read("STUDY_TEST", selected)
    assert content is not None
    assert len(content) > 0


def test_default_generation_flow_uses_builtin_template():
    draft = create_sample_draft()
    docx_bytes = docx_generation_service.generate_docx("NCT00081328", draft)
    text = document_text(docx_bytes)

    assert "Informed Consent Form" in text or "INFORMED CONSENT FORM" in text or "Information Sheet" in text
    assert "NOT_FOUND" not in text
    assert "Not provided in the study protocol." not in text


def test_unsupported_sections_removed_and_headings_renumbered():
    draft = create_sample_draft()
    docx_bytes = docx_generation_service.generate_docx("NCT00081328", draft)
    doc = docx.Document(io.BytesIO(docx_bytes))

    heading_texts = [p.text for p in doc.paragraphs if p.style and p.style.name and p.style.name.lower().startswith("heading 1")]

    # Alternatives section (NOT_FOUND) should be removed
    assert not any("Alternatives" in h for h in heading_texts)
    # Biological samples (no facts) should be removed
    assert not any("Biological Samples" in h for h in heading_texts)

    # Check sequential numbering with no gaps
    numbers = []
    for h in heading_texts:
        if h.split(".")[0].isdigit():
            numbers.append(int(h.split(".")[0]))

    assert numbers == list(range(1, len(numbers) + 1))


def test_signature_fields_preserved():
    draft = create_sample_draft()
    docx_bytes = docx_generation_service.generate_docx("NCT00081328", draft)
    text = document_text(docx_bytes)

    # Verify signature blocks exist in output
    assert "Signature of Participant" in text or "Participant Signature" in text or "PARTICIPANT_SIGNATURE" in text or "PARTICIPANT_NAME" in text
    assert "Signature" in text


def test_patient_friendly_template_rendering():
    draft = create_sample_draft()
    standard_bytes = docx_generation_service.generate_docx("NCT00081328", draft, export_type="standard")
    patient_bytes = docx_generation_service.generate_docx("NCT00081328", draft, export_type="patient_facing")

    doc_std = docx.Document(io.BytesIO(standard_bytes))
    doc_pf = docx.Document(io.BytesIO(patient_bytes))

    assert "Patient-Facing ICF" in doc_std.core_properties.title
    assert "Patient-Facing ICF" in doc_pf.core_properties.title
