import pytest
import io
import docx
from app.models.human_review import HumanReviewUpdateRequest, ReviewStatusEnum
from app.services.icf_generation_service import icf_generation_service
from app.services.template_service import template_service
from app.services.human_review_service import human_review_service
from app.models.icf import ICFGenerationRequest

def make_dummy_docx():
    doc = docx.Document()
    doc.add_heading('1. Purpose of the Study', level=1)
    doc.add_paragraph('{{icf.purpose}}')
    doc.add_heading('2. Voluntary Participation', level=1)
    doc.add_paragraph('{{icf.voluntary}}')
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def test_first_revision_builtin_template_approval_eligibility():
    study_id = 'TEST_REV1_BUILTIN'
    # Ensure clean slate
    template_service.remove(study_id)
    human_review_service.reset_reviews()
    
    # 1. Generate Revision 1
    req = ICFGenerationRequest(study_id=study_id, revision=1)
    resp = icf_generation_service.generate_icf(req)
    
    assert resp.revision == 1
    assert resp.icf_draft.revision == 1
    assert resp.template is not None
    assert resp.template.template_id == 'builtin-v1'
    assert resp.template.kind == 'default'
    
    # 2. Approve Revision 1 directly
    upd = HumanReviewUpdateRequest(review_status=ReviewStatusEnum.APPROVED, revision=1)
    review_state = human_review_service.update_review(study_id, upd)
    
    assert review_state.review_status == ReviewStatusEnum.APPROVED
    assert review_state.approved_revision == 1
    assert review_state.approved_template_id == 'builtin-v1'

def test_first_revision_custom_template_approval_eligibility():
    study_id = 'TEST_REV1_CUSTOM'
    template_service.remove(study_id)
    human_review_service.reset_reviews()
    
    # 1. Upload custom template
    docx_bytes = make_dummy_docx()
    selected = template_service.upload(study_id, 'custom_layout.docx', docx_bytes)
    assert selected.kind == 'custom'
    
    # 2. Generate Revision 1
    req = ICFGenerationRequest(study_id=study_id, revision=1)
    resp = icf_generation_service.generate_icf(req)
    
    assert resp.revision == 1
    assert resp.template.template_id == selected.template_id
    assert resp.template.kind == 'custom'
    
    # 3. Approve Revision 1 (with template_reviewed=True)
    upd = HumanReviewUpdateRequest(review_status=ReviewStatusEnum.APPROVED, revision=1, template_reviewed=True)
    review_state = human_review_service.update_review(study_id, upd)
    
    assert review_state.review_status == ReviewStatusEnum.APPROVED
    assert review_state.approved_revision == 1
    assert review_state.approved_template_id == selected.template_id

def test_multiple_different_studies_and_templates():
    study_a = 'TEST_MULTI_STUDY_A'
    study_b = 'TEST_MULTI_STUDY_B'
    template_service.remove(study_a)
    template_service.remove(study_b)
    
    docx_bytes = make_dummy_docx()
    sel_b = template_service.upload(study_b, 'b_layout.docx', docx_bytes)
    
    resp_a = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_a))
    resp_b = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_b))
    
    assert resp_a.template.kind == 'default'
    assert resp_b.template.kind == 'custom'
    assert resp_b.template.template_id == sel_b.template_id

def test_second_revision_preserves_existing_behavior():
    study_id = 'TEST_REV2_BEHAVIOR'
    template_service.remove(study_id)
    human_review_service.reset_reviews()
    
    # Rev 1
    resp1 = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_id))
    assert resp1.revision == 1
    
    # Rev 2 with feedback
    req2 = ICFGenerationRequest(study_id=study_id, reviewer_feedback='Simplify procedures section')
    resp2 = icf_generation_service.generate_icf(req2)
    assert resp2.revision == 2
    assert resp2.icf_draft.revision == 2
    assert resp2.template.template_id == resp1.template.template_id

def test_real_template_mismatch_blocked_backend():
    study_id = 'TEST_TEMPLATE_MISMATCH'
    template_service.remove(study_id)
    human_review_service.reset_reviews()
    
    # 1. Generate Rev 1 with default template
    resp1 = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_id))
    assert resp1.template.kind == 'default'
    
    # 2. User uploads NEW template AFTER Rev 1 generation
    docx_bytes = make_dummy_docx()
    sel_new = template_service.upload(study_id, 'new_template.docx', docx_bytes)
    
    # 3. Attempting to approve Rev 1 must fail due to template mismatch
    upd = HumanReviewUpdateRequest(review_status=ReviewStatusEnum.APPROVED, revision=1, template_reviewed=True)
    with pytest.raises(ValueError, match='selected template differs from this draft'):
        human_review_service.update_review(study_id, upd)

def test_missing_or_invalid_template_metadata_fails_safely():
    study_id = 'TEST_CORRUPT_TEMPLATE'
    template_service.remove(study_id)
    human_review_service.reset_reviews()
    
    resp = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_id))
    stored = icf_generation_service.get_stored_icf(study_id)
    stored.template = None
    
    upd = HumanReviewUpdateRequest(review_status=ReviewStatusEnum.APPROVED, revision=1)
    with pytest.raises(ValueError, match='selected template differs'):
        human_review_service.update_review(study_id, upd)

def test_backend_approval_enforcement_direct_call():
    study_id = 'TEST_BACKEND_ENFORCE'
    template_service.remove(study_id)
    human_review_service.reset_reviews()
    
    resp = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_id))
    
    # Approval without validation report matching revision must fail
    stored = icf_generation_service.get_stored_icf(study_id)
    stored.validation_report = {'revision': 99, 'overall_status': 'PASSED'}
    
    upd = HumanReviewUpdateRequest(review_status=ReviewStatusEnum.APPROVED, revision=1)
    with pytest.raises(ValueError, match='Validate the current revision'):
        human_review_service.update_review(study_id, upd)

def test_persistence_and_retrieval_across_revisions():
    study_id = 'TEST_PERSISTENCE'
    template_service.remove(study_id)
    
    resp1 = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_id))
    resp2 = icf_generation_service.generate_icf(ICFGenerationRequest(study_id=study_id, reviewer_feedback='Revise risks'))
    
    history = icf_generation_service.get_history(study_id)
    assert len(history) == 2
    assert history[0].revision == 1
    assert history[1].revision == 2
    assert history[0].template.template_id == history[1].template.template_id
