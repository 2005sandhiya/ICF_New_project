import io

import docx
from docx.shared import Inches, Pt

from app.models.icf import ICFDraft, ICFSection
from app.services.docx_generation_service import docx_generation_service
from app.services.template_renderer import document_text


def draft(content="Protocol-grounded purpose.", revision=2):
    return ICFDraft(revision=revision, sections=[
        ICFSection(section_id="purpose", title="Purpose of the Study", content=content, status="SUPPORTED", source_evidence_ids=["source-1"]),
        ICFSection(section_id="risks", title="Risks and Discomforts", content="Not provided in the study protocol."),
    ])


def save(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_split_placeholders_preserve_layout_tables_headers_and_styles():
    template = docx.Document()
    template.sections[0].left_margin = Inches(.7)
    template.sections[0].header.paragraphs[0].text = "Study {{study_id}} / Revision {{revision}}"
    template.sections[0].footer.paragraphs[0].text = "Consent office"
    template.add_heading("Our purpose", 1)
    paragraph = template.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(17)
    paragraph.add_run("{{icf.").bold = True
    paragraph.add_run("purpose}}")
    table = template.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Risks"
    table.cell(0, 1).text = "{{icf.risks}}"
    result = docx_generation_service.generate_docx("A", draft(), template_content=save(template))
    output = docx.Document(io.BytesIO(result))
    assert output.sections[0].left_margin == template.sections[0].left_margin
    assert output.sections[0].header.paragraphs[0].text == "Study A / Revision 2"
    assert output.sections[0].footer.paragraphs[0].text == "Consent office"
    assert "not provided in the study protocol" in output.tables[0].cell(0, 1).text
    populated = next(p for p in output.paragraphs if p.text == "Protocol-grounded purpose.")
    assert populated.runs[0].bold
    assert populated.paragraph_format.space_after == Pt(17)
    assert document_text(result).count("Protocol-grounded purpose.") == 1
    assert "{{" not in document_text(result)


def test_heading_population_replaces_sample_facts_and_appends_missing_sections():
    template = docx.Document()
    template.add_heading("Purpose of the Study", 1)
    template.add_paragraph("SAMPLE: Every participant receives 999 mg of invented drug.")
    template.add_paragraph("SAMPLE: Guaranteed cure.")
    template.add_heading("Signature instructions", 1)
    template.add_paragraph("Sign after your questions have been answered.")
    result = docx_generation_service.generate_docx("A", draft(), template_content=save(template))
    text = document_text(result)
    assert "999" not in text and "Guaranteed cure" not in text
    assert "Sign after your questions have been answered." in text
    assert text.count("Protocol-grounded purpose.") == 1


def test_table_heading_maps_to_adjacent_cell():
    template = docx.Document()
    table = template.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Purpose of the Study"
    table.cell(0, 1).text = "Old study sample."
    result = docx_generation_service.generate_docx("A", draft(), template_content=save(template))
    output = docx.Document(io.BytesIO(result))
    assert output.tables[0].cell(0, 1).text == "Protocol-grounded purpose."
    assert "Old study sample" not in document_text(result)


def test_standard_and_patient_friendly_use_their_own_content():
    template = docx.Document()
    template.add_paragraph("{{icf.purpose}}")
    content = save(template)
    standard = docx_generation_service.generate_docx("A", draft("Evaluate treatment efficacy."), template_content=content)
    patient = docx_generation_service.generate_docx("A", draft("Check how well the treatment works."), export_type="patient_friendly", template_content=content)
    assert "Evaluate treatment efficacy." in document_text(standard)
    assert "Check how well the treatment works." in document_text(patient)
    assert "Evaluate treatment efficacy." not in document_text(patient)


def test_default_layout_remains_available():
    result = docx_generation_service.generate_docx("A", draft())
    text = document_text(result)
    assert "INFORMED CONSENT FORM" in text
    assert "Revision 2" in text
    assert "Protocol-grounded purpose." in text
