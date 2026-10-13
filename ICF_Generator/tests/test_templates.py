import io
from zipfile import ZIP_DEFLATED, ZipFile

import docx
import pytest

from app.services.template_service import MAX_TEMPLATE_BYTES, TemplateError, TemplateService


def template_bytes(text="{{icf.purpose}}"):
    document = docx.Document()
    document.add_heading("Purpose of the Study", 1)
    document.add_paragraph(text)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.fixture
def templates(tmp_path):
    return TemplateService(tmp_path)


def test_default_without_upload(templates):
    selected = templates.selection("STUDY_A")
    assert selected.kind == "default"
    content = templates.read("STUDY_A", selected)
    assert content is not None
    assert len(content) > 0


def test_valid_upload_and_safe_filename(templates):
    content = template_bytes()
    selected = templates.upload("../STUDY_A", "../../sample.docx", content)
    assert selected.filename == "sample.docx"
    assert selected.mapped_sections == ["purpose"]
    assert "risks" in selected.missing_sections
    assert templates.read("../STUDY_A", selected) == content
    assert templates.selection("../STUDY_A") == selected
    assert len(list(templates.root.glob("*/selection.json"))) == 1


@pytest.mark.parametrize("filename,content", [
    ("bad.pdf", b"%PDF-1.4"), ("bad.doc", b"legacy"),
    ("bad.docx", b"broken zip"), ("empty.docx", b""),
    ("large.docx", b"a" * (MAX_TEMPLATE_BYTES + 1)),
], ids=["pdf", "legacy-doc", "corrupt", "empty", "oversized"])
def test_invalid_upload_does_not_change_selection(templates, filename, content):
    good = templates.upload("A", "good.docx", template_bytes())
    with pytest.raises(TemplateError):
        templates.upload("A", filename, content)
    assert templates.selection("A") == good


def test_replace_remove_and_study_isolation(templates):
    original = template_bytes()
    first = templates.upload("A", "first.docx", original)
    second = templates.upload("A", "second.docx", template_bytes("{{icf.risks}}"))
    assert first.template_id != second.template_id
    assert templates.selection("A") == second
    assert templates.selection("B").kind == "default"
    with pytest.raises(TemplateError, match="study mismatch"):
        templates.read("B", second)
    assert templates.remove("A").kind == "default"
    assert templates.read("A", first) == original


def test_missing_or_modified_custom_template_never_falls_back(templates):
    selected = templates.upload("A", "one.docx", template_bytes())
    stored = templates._folder("A") / (selected.template_id + ".docx")
    stored.write_bytes(b"changed")
    with pytest.raises(TemplateError, match="selected template"):
        templates.selection("A")
    stored.unlink()
    with pytest.raises(TemplateError, match="selected template"):
        templates.selection("A")
    assert templates.remove("A").kind == "default"


@pytest.mark.parametrize("text", ["{{icf.unknown}}", "{{icf.purpose}} {{icf.purpose}}"])
def test_ambiguous_or_unknown_slots_rejected(templates, text):
    with pytest.raises(TemplateError):
        templates.upload("A", "bad.docx", template_bytes(text))


@pytest.mark.parametrize("member,value", [
    ("../bad.xml", b"<a/>"), ("word/vbaProject.bin", b"code"),
    ("word/embeddings/oleObject1.bin", b"code"),
    ("word/evil.xml", b'<!DOCTYPE foo [<!ENTITY x "bad">]><foo>&x;</foo>'),
    ("word/_rels/evil.xml.rels", b'<Relationships><Relationship TargetMode="External" Type="image" Target="file:///secret"/></Relationships>'),
])
def test_unsafe_packages_rejected(templates, member, value):
    stream = io.BytesIO(template_bytes())
    with ZipFile(stream, "a", ZIP_DEFLATED) as archive:
        archive.writestr(member, value)
    with pytest.raises(TemplateError):
        templates.upload("A", "unsafe.docx", stream.getvalue())
